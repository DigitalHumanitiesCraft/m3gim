"""Source-fidelity contract (knowledge/data.md, E-301)."""

from __future__ import annotations

from collections import defaultdict
import hashlib
from pathlib import Path
import sys

import pandas as pd
from openpyxl import load_workbook
from rdflib import Graph, Literal, Namespace, RDFS

from _helpers import ensure_list

REPO_ROOT = Path(__file__).parent.parent
sys.path.insert(0, str(REPO_ROOT / "scripts"))
from _common import link_record_key, resolve_record_key


M3 = "m3gim-ontology:"
RECORDED = (f"{M3}recordedType", f"{M3}recordedValue", f"{M3}recordedRole")
LEGACY_COMPONENT_PROPERTIES = ("hasPerformance", "hasPerformer", "performanceOf", "hasStageRole")
INDEX_PROPERTIES = {
    "composer",
    f"{M3}headquarters",
    f"{M3}indexNote",
    f"{M3}keyContact",
    f"{M3}lifespan",
    f"{M3}sungPart",
}
WIKIDATA_PROPERTIES = {
    "gndo:professionOrOccupationAsLiteral",
    "geo:lat",
    "geo:long",
    "schema:birthDate",
    "schema:birthPlace",
    "schema:deathDate",
    "schema:deathPlace",
    f"{M3}country",
    f"{M3}voiceType",
    f"{M3}wdComposer",
    f"{M3}wdGenre",
    f"{M3}wdInception",
    f"{M3}wdLocation",
    f"{M3}wdPremiereDate",
    f"{M3}wdPublicationDate",
}


def _walk(value):
    if isinstance(value, dict):
        yield value
        for child in value.values():
            yield from _walk(child)
    elif isinstance(value, list):
        for child in value:
            yield from _walk(child)


def _text(value) -> str:
    return "" if pd.isna(value) else str(value).strip()


def _raw(value):
    return "" if pd.isna(value) else value


def _source_key(node):
    source = node.get(f"{M3}xlsxSource")
    if not isinstance(source, dict):
        return None
    return (source.get(f"{M3}xlsxSheet"), source.get(f"{M3}xlsxRow"))


def _resolvable_link_rows(frame, graph):
    known = {
        node.get("rico:identifier")
        for node in graph
        if node.get("rico:identifier")
    }
    for _, row in frame.iterrows():
        signature = _text(row.get("archivsignatur"))
        if not signature:
            continue
        key = link_record_key(signature, _text(row.get("folio")))
        if resolve_record_key(key, known):
            yield row


def test_legacy_statement_paths_have_no_event_or_participation_entailment():
    graph = Graph().parse(REPO_ROOT / "vocab/m3gim.ttl", format="turtle")
    m3 = Namespace("https://dhcraft.org/m3gim/ontology#")

    for local_name in LEGACY_COMPONENT_PROPERTIES:
        prop = m3[local_name]
        domains = list(graph.objects(prop, RDFS.domain))
        if local_name != "hasPerformance":
            assert not domains, local_name
        assert not list(graph.objects(prop, RDFS.range)), local_name
        assert not list(graph.objects(prop, RDFS.subPropertyOf)), local_name

    for local_name in ("atDate", "atPlace"):
        assert not list(graph.objects(m3[local_name], RDFS.subPropertyOf)), local_name


def test_recorded_link_cells_round_trip_on_neutral_carriers(graph, xlsx_verknuepfungen):
    carriers = defaultdict(list)
    for node in _walk(graph):
        if any(field in node for field in RECORDED):
            carriers[_source_key(node)].append(node)

    missing = []
    mismatched = []
    for row in _resolvable_link_rows(xlsx_verknuepfungen, graph):
        recorded_type = _raw(row.get("typ"))
        recorded_value = _raw(row.get("name"))
        key = (_text(row.get("_xlsx_sheet")), int(row.get("_xlsx_row")))
        candidates = carriers.get(key, [])
        if not any(_text(row.get(field)) for field in ("name", "rolle", "datum", "anmerkung")):
            assert not candidates, f"Incomplete row became a statement: {key}"
            continue
        if not candidates:
            missing.append(key)
            continue
        expected = (recorded_type, recorded_value, _raw(row.get("rolle")))
        actual = {
            tuple(_raw(node.get(field)) for field in RECORDED)
            for node in candidates
        }
        if expected not in actual:
            mismatched.append((key, expected, sorted(actual)))

    assert not missing, f"Quellzeilen ohne neutralen Aussageknoten: {missing[:10]}"
    assert not mismatched, f"Nicht wortgetreu serialisierte Quellzellen: {mismatched[:5]}"


def test_recorded_notes_remain_on_their_source_carrier(graph, xlsx_verknuepfungen):
    by_source = defaultdict(list)
    for node in _walk(graph):
        if any(field in node for field in RECORDED):
            by_source[_source_key(node)].append(node)

    offenders = []
    for row in _resolvable_link_rows(xlsx_verknuepfungen, graph):
        note = _raw(row.get("anmerkung"))
        if not note:
            continue
        key = (_text(row.get("_xlsx_sheet")), int(row.get("_xlsx_row")))
        if not any(_raw(n.get("rico:generalDescription")) == note for n in by_source.get(key, [])):
            offenders.append((key, note))
    assert not offenders, f"Quellanmerkungen fehlen oder wurden verändert: {offenders[:5]}"


def test_dating_evidence_round_trips_by_object_source_row(graph, xlsx_objekte):
    by_source = {
        _source_key(node): node
        for node in graph
        if node.get("@type") in {"rico:Record", "rico:RecordSet"}
        and _source_key(node)
    }
    offenders = []
    for index, row in xlsx_objekte.iterrows():
        evidence = _text(row.get("datierungsevidenz"))
        if not evidence:
            continue
        key = ("Objekte", int(index) + 2)
        actual = _text(by_source.get(key, {}).get(f"{M3}datingEvidence"))
        if actual != evidence:
            offenders.append((key, evidence, actual))
    assert not offenders, f"Datierungsevidenz fehlt oder wurde verändert: {offenders[:5]}"


def test_every_added_property_has_matching_property_source(graph):
    offenders = []
    for node in _walk(graph):
        present = (INDEX_PROPERTIES | WIKIDATA_PROPERTIES).intersection(node)
        if not present:
            continue
        descriptors = ensure_list(node.get(f"{M3}propertySource"))
        for prop in present:
            kind = "index" if prop in INDEX_PROPERTIES else "wikidata"
            matches = [
                source for source in descriptors
                if isinstance(source, dict)
                and source.get(f"{M3}sourceProperty") == prop
                and source.get(f"{M3}sourceKind") == kind
            ]
            if not matches:
                offenders.append((node.get("@id") or node.get("name"), prop, kind))
                continue
            for source in matches:
                assert source.get(f"{M3}sourceValue") not in (None, "")
                assert source.get("dcterms:source")
                if kind == "index":
                    assert isinstance(source.get(f"{M3}xlsxSource"), dict)
    assert not offenders, f"Zusatzwerte ohne eigenschaftsspezifische Quelle: {offenders[:10]}"


def test_property_source_values_survive_rdf_expansion(jsonld_path):
    rdf = Graph().parse(jsonld_path, format="json-ld")
    m3 = Namespace("https://dhcraft.org/m3gim/ontology#")
    descriptors = set(rdf.objects(None, m3.propertySource))
    assert descriptors, "Keine Eigenschaftsquellen im RDF-Graph"
    missing = [node for node in descriptors if not list(rdf.objects(node, m3.sourceValue))]
    non_literals = [
        (node, value)
        for node in descriptors
        for value in rdf.objects(node, m3.sourceValue)
        if not isinstance(value, Literal)
    ]
    assert not missing, f"sourceValue ging bei JSON-LD-Expansion verloren: {missing[:5]}"
    assert not non_literals, f"sourceValue expandiert als kontextloser Knoten: {non_literals[:5]}"


def _authority_source_value(properties, source_property):
    paths = {
        "gndo:professionOrOccupationAsLiteral": ("occupation", "labels"),
        f"{M3}voiceType": ("voiceType", "first_label"),
        "schema:birthDate": ("birthDate", "raw"),
        "schema:deathDate": ("deathDate", "raw"),
        "schema:birthPlace": ("birthPlace", "label"),
        "schema:deathPlace": ("deathPlace", "label"),
        "geo:lat": ("coordinates", "lat"),
        "geo:long": ("coordinates", "lon"),
        f"{M3}country": ("country", "label"),
        f"{M3}wdComposer": ("composer", "label"),
        f"{M3}wdGenre": ("genre", "labels"),
        f"{M3}wdPremiereDate": ("premiereDate", "raw"),
        f"{M3}wdPublicationDate": ("publicationDate", "raw"),
        f"{M3}wdLocation": ("location", "label"),
        f"{M3}wdInception": ("inception", "raw"),
    }
    source_key, selector = paths[source_property]
    raw = properties[source_key]
    if selector == "raw":
        return raw
    if selector in {"lat", "lon"}:
        return raw[selector]
    if selector == "label":
        return raw.get("label", raw.get("qid", ""))
    if selector == "first_label":
        item = raw[0] if isinstance(raw, list) else raw
        return item.get("label", "") if isinstance(item, dict) else str(item)
    return [item.get("label", item.get("qid", "")) for item in ensure_list(raw) if isinstance(item, dict)]


def test_wikidata_source_values_match_the_property_path(graph, enrichment):
    assert enrichment, "Wikidata-Cache für den Herkunftsabgleich fehlt"
    entities = enrichment.get("entities", enrichment)
    mismatches = []
    for node in _walk(graph):
        for descriptor in ensure_list(node.get(f"{M3}propertySource")):
            if not isinstance(descriptor, dict) or descriptor.get(f"{M3}sourceKind") != "wikidata":
                continue
            prop = descriptor.get(f"{M3}sourceProperty")
            if prop not in WIKIDATA_PROPERTIES:
                continue
            source = descriptor.get("dcterms:source", {})
            uri = source.get("@id", "") if isinstance(source, dict) else str(source)
            qid = uri.rsplit("/", 1)[-1]
            properties = entities.get(qid, {}).get("properties", {})
            if not properties:
                mismatches.append((qid, prop, "Cache-Eintrag fehlt"))
                continue
            expected = _authority_source_value(properties, prop)
            actual = descriptor.get(f"{M3}sourceValue")
            differs = (
                actual != expected and actual not in expected
                if isinstance(expected, list) else actual != expected
            )
            if differs:
                mismatches.append((qid, prop, expected, actual))
    assert not mismatches, f"Wikidata-sourceValue entspricht nicht seinem Quellpfad: {mismatches[:5]}"


def test_index_source_values_exist_in_the_declared_source_row(graph, sheets_dir):
    workbooks = {}
    paths_by_source = {
        f"urn:sha256:{hashlib.sha256(path.read_bytes()).hexdigest()}": path
        for path in sheets_dir.glob("M3GIM-*index.xlsx")
    }
    assert len(paths_by_source) == 4, "Die vier Indexquellen wurden nicht eindeutig gefunden"
    mismatches = []
    for node in _walk(graph):
        for descriptor in ensure_list(node.get(f"{M3}propertySource")):
            if not isinstance(descriptor, dict) or descriptor.get(f"{M3}sourceKind") != "index":
                continue
            source = descriptor.get("dcterms:source", {})
            uri = source.get("@id", "") if isinstance(source, dict) else str(source)
            path = paths_by_source.get(uri)
            address = descriptor.get(f"{M3}xlsxSource", {})
            row_number = address.get(f"{M3}xlsxRow") if isinstance(address, dict) else None
            if not path or not path.exists() or not isinstance(row_number, int):
                mismatches.append((uri, row_number, "Quelle nicht lesbar"))
                continue
            if path not in workbooks:
                workbooks[path] = load_workbook(path, read_only=True, data_only=True)
            workbook = workbooks[path]
            sheet_name = address.get(f"{M3}xlsxSheet")
            if sheet_name not in workbook.sheetnames:
                mismatches.append((path.name, sheet_name, "Blatt fehlt"))
                continue
            sheet = workbook[sheet_name]
            row_values = [_raw(cell.value) for cell in sheet[row_number] if _raw(cell.value) != ""]
            value = descriptor.get(f"{M3}sourceValue")
            if isinstance(value, (dict, list)) or value not in row_values:
                mismatches.append((path.name, row_number, value, [repr(item) for item in row_values]))
    assert not mismatches, f"Index-sourceValue ist nicht der Wert seiner angegebenen Zeile: {mismatches[:5]}"


def test_authority_ids_do_not_carry_document_local_context(graph):
    local_properties = {
        "role",
        f"{M3}derivedFromRole",
        f"{M3}recordedRole",
        f"{M3}recordedType",
        f"{M3}recordedValue",
        f"{M3}xlsxSource",
    }
    offenders = []
    for node in _walk(graph):
        identifier = node.get("@id", "")
        leaked = sorted(local_properties.intersection(node))
        if identifier.startswith("wd:Q") and leaked:
            offenders.append((identifier, leaked))
    assert not offenders, (
        "Dokumentlokale Aussage an globaler Wikidata-Identität; RDF würde Kontexte "
        f"verschmelzen: {offenders[:10]}"
    )


def test_authority_identity_is_a_sourced_reference_without_owl_identity(graph):
    offenders = []
    seen = 0
    for node in _walk(graph):
        reference = node.get(f"{M3}authorityReference")
        if not isinstance(reference, dict):
            continue
        seen += 1
        descriptors = ensure_list(node.get(f"{M3}propertySource"))
        matching = [
            source for source in descriptors
            if isinstance(source, dict)
            and source.get(f"{M3}sourceProperty") == f"{M3}authorityReference"
            and source.get(f"{M3}sourceKind") in {"index", "wikidata"}
        ]
        if (
            not str(reference.get("@id", "")).startswith("wd:Q")
            or "@id" in node
            or "owl:sameAs" in node
            or not matching
        ):
            offenders.append((node.get("name"), reference, bool(matching)))
    assert seen, "Keine separat belegte Normdatenreferenz im Korpus"
    assert not offenders, f"Normdatenidentität nicht lokal und belegtreu modelliert: {offenders[:10]}"


def test_source_mentions_do_not_create_performances_or_name_global_stage_roles(graph):
    performances = [node.get("@id") for node in graph if node.get("@type") == f"{M3}Performance"]
    assert not performances, f"Quellgruppierung als Aufführung serialisiert: {performances[:10]}"

    roles_by_name = defaultdict(list)
    for node in graph:
        if node.get("@type") != f"{M3}StageRole":
            continue
        source = _source_key(node)
        name = _text(node.get("rico:name") or node.get("name") or node.get(f"{M3}recordedValue")).casefold()
        assert source, f"Bühnenrollen-Nennung ohne lokale Quellzeile: {node.get('@id')}"
        assert name, f"Bühnenrollen-Nennung ohne erfassten Wert: {node.get('@id')}"
        roles_by_name[name].append((node.get("@id"), source))

    collapsed = []
    for name, occurrences in roles_by_name.items():
        source_count = len({source for _, source in occurrences})
        id_count = len({identifier for identifier, _ in occurrences})
        if id_count < source_count:
            collapsed.append((name, occurrences))
    assert not collapsed, f"Gleichnamige Partien über Quellzeilen identifiziert: {collapsed[:5]}"
